"""
=============================================================
Model 2 — Protocol Tolerance Prediction (Multi-class)
"What is the strictest QA protocol this plan can pass?"


  Class 0 — FAIL ALL protocols           (' Fail QA' sheet)
  Class 1 — FAIL dosisoft, PASS lenient  (' GPR' sheet)
  Class 2 — PASS dosisoft (strictest)    (' GPR' sheet)


Clinical logic: dosisoft (2%/2mm) is always tried first.
A plan whose only recorded measurement is a lenient protocol
therefore failed dosisoft.

Features: 12 complexity metrics + anatomical site

Models: Random Forest + XGBoost
=============================================================
"""

# ================================================================
#  IMPORTANT — NOT FOR GENERAL USE
# ================================================================
# This script is provided for REFERENCE and reproducibility of the
# results presented in this thesis. It requires the labelled dataset
# 'Per_protocol.xlsx' which contains patient-related clinical data
# and CANNOT be shared due to confidentiality.
#
# This script CANNOT be executed without the corresponding dataset.
# It is NOT intended as a general-purpose tool for other datasets.
#
# For reusable code, see:
#   - all_metrics.py   (VMAT complexity extraction from DICOM-RT)
#   - csv_to_excel.py  (CSV to Excel converter)
# ================================================================

import pandas as pd
import numpy as np
import matplotlib.pyplot as plt
from sklearn.ensemble import RandomForestClassifier
from sklearn.model_selection import StratifiedKFold, cross_validate, cross_val_predict
from sklearn.metrics import confusion_matrix, classification_report, ConfusionMatrixDisplay
from sklearn.preprocessing import LabelEncoder
import joblib
import openpyxl
import warnings
warnings.filterwarnings('ignore')

try:
    from xgboost import XGBClassifier
    HAS_XGB = True
except ImportError:
    print("  XGBoost not found. Install with: pip install xgboost")
    HAS_XGB = False


# ============================================================
#  USER SETTINGS — CHANGE THIS PATH FOR YOUR SYSTEM
# ============================================================

FILE = "Per_protocol.xlsx"  # Change to full path if needed

# ============================================================


# =============================================================
# 1. DATA LOADING & 3-CLASS LABEL CONSTRUCTION
# =============================================================
print("=" * 60)
print(" MODEL 2 — PROTOCOL TOLERANCE PREDICTION (3-class)")
print("=" * 60)

METRIC_COLS = [
    'MCSv', 'Total_MU', 'LSV_Mean', 'AAV_Mean', 'MI_Mean',
    'Leaf_Travel_mm', 'LTMCS', 'MU_per_CP_Mean', 'MU_lt_3_Percent',
    'SAS5_Mean', 'SAS10_Mean', 'AAJA_Mean'
]

def norm_proto(p):
    if not p or str(p).strip() in ('', 'nan'): return 'Unknown'
    p = str(p).replace('\xa0', ' ').lower().strip()
    if 'dosisoft' in p: return 'dosisoft'
    if '90%' in p:      return '3mm90'
    if '3% 2mm' in p or '3%2mm' in p: return '3mm2'
    if '3% 3mm' in p or '3%3mm' in p: return '3mm3'
    return 'Unknown'

def norm_site(d):
    if not d: return 'Unknown'
    d = str(d).upper()
    if 'PROSTATE' in d:                                                return 'Prostate'
    if any(x in d for x in ['ENDOMETRIUM','CERVIX','VULVA','OVARY']):  return 'Gynecological'
    if any(x in d for x in ['BLADDER','RECTUM','ANUS']):               return 'Pelvic'
    return 'Other'

# Robust sheet lookup (handles inconsistent leading/trailing spaces)
wb = openpyxl.load_workbook(FILE, read_only=True)
sheet_map = {s.strip().lower(): s for s in wb.sheetnames}

def find_sheet(target):
    key = target.strip().lower()
    if key not in sheet_map:
        raise ValueError(f"Could not find sheet matching '{target}'. Sheets: {wb.sheetnames}")
    return sheet_map[key]

GPR_SHEET  = find_sheet('GPR')
FAIL_SHEET = find_sheet('Fail QA')

# ── GPR sheet: build Class 1 / Class 2 (one row per UNIQUE plan) ──
df_gpr = pd.read_excel(FILE, sheet_name=GPR_SHEET)
df_gpr['Protocol'] = df_gpr['Πρωτόκολλο'].apply(norm_proto)
df_gpr['Site']     = df_gpr['Διάγνωση'].apply(norm_site)

# Clinical rule: dosisoft is always tried first.
#   -> passed_dosisoft == True  for any measurement row  => Class 2
#   -> otherwise (only lenient protocol(s) recorded)      => Class 1
records = []
for pid, group in df_gpr.groupby('ID'):
    passed_dosisoft = ((group['Status'] == 'ΕΠΙΤΥΧΙΑ') & (group['Protocol'] == 'dosisoft')).any()
    qa_class = 2 if passed_dosisoft else 1

    pass_row = group[group['Status'] == 'ΕΠΙΤΥΧΙΑ'].iloc[0]
    rec = {m: pass_row[m] for m in METRIC_COLS}
    rec['Site'] = pass_row['Site']
    rec['QA_Class'] = qa_class
    records.append(rec)

df_gpr_plans = pd.DataFrame(records)

# ── Fail QA sheet: Class 0 (failed ALL protocols) ──────────
df_fail = pd.read_excel(FILE, sheet_name=FAIL_SHEET)
df_fail['Site'] = df_fail['Διάγνωση (filename)'].apply(norm_site)
df_fail['QA_Class'] = 0
df_fail_plans = df_fail[METRIC_COLS + ['Site', 'QA_Class']]

# ── Combine ──────────────────────────────────────────────────
df = pd.concat([df_gpr_plans, df_fail_plans], ignore_index=True)
df = df.dropna(subset=METRIC_COLS)

class_names = {0: 'FAIL ALL', 1: 'FAIL strict / PASS lenient', 2: 'PASS strict (dosisoft)'}

print(f"\nClass distribution (one row per plan, N={len(df)}):")
for c in [0, 1, 2]:
    n = (df['QA_Class'] == c).sum()
    print(f"  Class {c} ({class_names[c]}): {n}")

# Encode site
le_site = LabelEncoder()
df['Site_enc'] = le_site.fit_transform(df['Site'].astype(str))
print(f"\n  Sites: {list(le_site.classes_)}")

FEATURE_COLS = METRIC_COLS + ['Site_enc']
X = df[FEATURE_COLS].fillna(df[FEATURE_COLS].median(numeric_only=True))
y = df['QA_Class']


# =============================================================
# 2. MODEL DEFINITIONS
# =============================================================
print("\n" + "=" * 60)
print(" MODEL DEFINITIONS")
print("=" * 60)

models = {
    "Random Forest": RandomForestClassifier(
        n_estimators=300,
        class_weight='balanced',
        min_samples_split=5,
        min_samples_leaf=2,
        max_features='sqrt',
        random_state=42,
        n_jobs=-1
    ),
}
if HAS_XGB:
    models["XGBoost"] = XGBClassifier(
        n_estimators=300,
        learning_rate=0.05,
        max_depth=4,
        subsample=0.8,
        colsample_bytree=0.8,
        objective='multi:softprob',
        num_class=3,
        random_state=42,
        verbosity=0,
        eval_metric='mlogloss'
    )

for name in models:
    print(f"  ✓ {name}")


# =============================================================
# 3. STRATIFIED 5-FOLD CROSS VALIDATION
# =============================================================
print("\n" + "=" * 60)
print(" STRATIFIED 5-FOLD CROSS VALIDATION")
print("=" * 60)

skf = StratifiedKFold(n_splits=5, shuffle=True, random_state=42)
results = {}
cm_data = {}

for name, model in models.items():
    cv = cross_validate(
        model, X, y, cv=skf,
        scoring=['accuracy', 'f1_macro', 'precision_macro', 'recall_macro'],
        return_train_score=False
    )
    results[name] = cv

    y_pred = cross_val_predict(model, X, y, cv=skf)
    cm = confusion_matrix(y, y_pred)
    cm_data[name] = cm

    print(f"\n{'─'*50}")
    print(f" {name}")
    print(f"{'─'*50}")
    print(f"  Accuracy       : {cv['test_accuracy'].mean():.3f} ± {cv['test_accuracy'].std():.3f}")
    print(f"  F1 (macro)     : {cv['test_f1_macro'].mean():.3f} ± {cv['test_f1_macro'].std():.3f}")
    print(f"  Precision (m)  : {cv['test_precision_macro'].mean():.3f} ± {cv['test_precision_macro'].std():.3f}")
    print(f"  Recall (macro) : {cv['test_recall_macro'].mean():.3f} ± {cv['test_recall_macro'].std():.3f}")
    print(f"\n  Classification report (cross-validated predictions):")
    print(classification_report(y, y_pred, target_names=[class_names[c] for c in [0,1,2]], digits=3))


# =============================================================
# 4. COMPARISON TABLE
# =============================================================
print("\n" + "=" * 60)
print(" RESULTS COMPARISON TABLE")
print("=" * 60)

summary = pd.DataFrame({
    name: {
        "Accuracy":       f"{v['test_accuracy'].mean():.3f}",
        "Accuracy (±std)":f"±{v['test_accuracy'].std():.3f}",
        "F1 (macro)":     f"{v['test_f1_macro'].mean():.3f}",
        "Precision (m)":  f"{v['test_precision_macro'].mean():.3f}",
        "Recall (macro)": f"{v['test_recall_macro'].mean():.3f}",
    }
    for name, v in results.items()
}).T
print(summary.to_string())
best = max(results, key=lambda n: results[n]['test_accuracy'].mean())
print(f"\n Best model (Accuracy): {best}")


# =============================================================
# 5. FEATURE IMPORTANCE + SAVE MODELS
# =============================================================
print("\n" + "=" * 60)
print(" FEATURE IMPORTANCE")
print("=" * 60)

importances = {}
trained_models = {}

rf_full = RandomForestClassifier(n_estimators=300, class_weight='balanced',
                                  min_samples_split=5, min_samples_leaf=2,
                                  max_features='sqrt', random_state=42, n_jobs=-1)
rf_full.fit(X, y)
importances["Random Forest"] = pd.Series(rf_full.feature_importances_,
                                          index=FEATURE_COLS).sort_values(ascending=False)
trained_models["Random Forest"] = rf_full

if HAS_XGB:
    xgb_full = XGBClassifier(n_estimators=300, learning_rate=0.05, max_depth=4,
                              subsample=0.8, colsample_bytree=0.8,
                              objective='multi:softprob', num_class=3,
                              random_state=42, verbosity=0, eval_metric='mlogloss')
    xgb_full.fit(X, y)
    importances["XGBoost"] = pd.Series(xgb_full.feature_importances_,
                                        index=FEATURE_COLS).sort_values(ascending=False)
    trained_models["XGBoost"] = xgb_full

for mname, imp in importances.items():
    print(f"\n{mname} — Top 8:")
    for i, (feat, val) in enumerate(imp.head(8).items(), 1):
        print(f"  {i}. {feat:<25} {val:.4f}")

for mname, model in trained_models.items():
    fname = f"model2_{'RF' if 'Forest' in mname else 'XGB'}.pkl"
    joblib.dump({
        'model': model,
        'feature_cols': FEATURE_COLS,
        'metric_cols': METRIC_COLS,
        'le_site': le_site,
        'class_names': class_names
    }, fname)
    print(f"\n  ✓ Saved: {fname}")


# =============================================================
# 6. PREDICTION FUNCTION
# =============================================================
def predict_protocol_tolerance(metrics_dict, site, model_file='model2_RF.pkl'):
    """
    Predict the strictest QA protocol a new plan can tolerate.

    Parameters:
    -----------
    metrics_dict : dict — complexity metrics from all_metrics.py output
    site         : str  — 'Prostate' | 'Gynecological' | 'Pelvic'
    model_file   : str  — path to saved model .pkl

    Returns:
    --------
    dict with predicted class, probabilities, and recommendation
    """
    saved = joblib.load(model_file)
    model    = saved['model']
    le_site  = saved['le_site']
    feat_cols= saved['feature_cols']
    met_cols = saved['metric_cols']
    cnames   = saved['class_names']

    site_enc = le_site.transform([site])[0] if site in le_site.classes_ else 0
    row = [metrics_dict.get(m, np.nan) for m in met_cols] + [site_enc]
    X_new = pd.DataFrame([row], columns=feat_cols)
    X_new = X_new.fillna(X_new.median(numeric_only=True))

    pred_class = int(model.predict(X_new)[0])
    probs = model.predict_proba(X_new)[0]

    recommendations = {
        0: "⚠️  Plan likely to FAIL QA under ALL protocols — re-planning strongly recommended.",
        1: "🟡 Plan likely to FAIL strict protocol (dosisoft) but PASS a lenient protocol. Recommend lenient protocol for QA.",
        2: "✅ Plan likely to PASS even the strictest protocol (dosisoft 2%/2mm)."
    }

    return {
        'predicted_class': pred_class,
        'class_label': cnames[pred_class],
        'probabilities': {cnames[i]: round(p*100, 1) for i, p in enumerate(probs)},
        'recommendation': recommendations[pred_class]
    }

print("\n✅ predict_protocol_tolerance() function defined")


# =============================================================
# 7. PLOTS
# =============================================================
print("\n" + "=" * 60)
print(" GENERATING PLOTS")
print("=" * 60)

fig, axes = plt.subplots(2, 3, figsize=(16, 10))
fig.suptitle("Model 2 — Protocol Tolerance Prediction (3-class)\n"
             f"(N={len(df)}: Class0={int((y==0).sum())}, "
             f"Class1={int((y==1).sum())}, Class2={int((y==2).sum())})",
             fontsize=13, fontweight='bold')

cm_labels = ['FAIL\nALL', 'FAIL strict\nPASS lenient', 'PASS\nstrict']

# Confusion matrices
for i, (name, cm) in enumerate(cm_data.items()):
    ax = axes[0, i]
    disp = ConfusionMatrixDisplay(confusion_matrix=cm, display_labels=cm_labels)
    disp.plot(ax=ax, cmap='Blues', colorbar=False, values_format='d')
    ax.set_title(f"{name} — Confusion Matrix")
    ax.tick_params(axis='x', labelsize=8)
    ax.tick_params(axis='y', labelsize=8)

# Accuracy comparison
ax3 = axes[0, 2]
acc_data = [results[n]['test_accuracy'] for n in results]
colors = ["#55A868", "#4C72B0"]
bp = ax3.boxplot(acc_data, patch_artist=True, widths=0.5)
for patch, color in zip(bp['boxes'], colors[:len(results)]):
    patch.set_facecolor(color)
    patch.set_alpha(0.75)
ax3.set_xticks(range(1, len(results)+1))
ax3.set_xticklabels(list(results.keys()), rotation=15, ha='right', fontsize=9)
ax3.set_ylabel("Accuracy")
ax3.set_title("Accuracy per Fold (5-CV)")
ax3.axhline(1/3, color='gray', linestyle='--', alpha=0.5, label='Random (1/3)')
ax3.legend(fontsize=8)
ax3.set_ylim(0, 1)

# Feature importance
ax4 = axes[1, 0:2]
top_feats = importances["Random Forest"].head(10)
bars = ax4.barh(top_feats.index[::-1], top_feats.values[::-1],
                color="#55A868", alpha=0.8, edgecolor='white')
ax4.set_xlabel("Feature Importance")
ax4.set_title("Top 10 Features — Random Forest")
for bar, val in zip(bars, top_feats.values[::-1]):
    ax4.text(bar.get_width()+0.002, bar.get_y()+bar.get_height()/2,
             f'{val:.3f}', va='center', fontsize=9)

# Class distribution
ax5 = axes[1, 2]
counts = [int((y==c).sum()) for c in [0,1,2]]
clrs = ['#C44E52', '#FF8C00', '#55A868']
bars5 = ax5.bar(cm_labels, counts, color=clrs, alpha=0.8, edgecolor='white')
for bar, val in zip(bars5, counts):
    ax5.text(bar.get_x()+bar.get_width()/2, bar.get_height()+1,
             str(val), ha='center', fontweight='bold')
ax5.set_ylabel("Number of plans")
ax5.set_title(f"Class Distribution (N={len(df)})")

plt.tight_layout()
plt.savefig("model2_protocol_tolerance.png", dpi=150, bbox_inches='tight')
plt.show()
print("\n✅ Plot saved as 'model2_protocol_tolerance.png'")
print("✅ Model 2 complete!")