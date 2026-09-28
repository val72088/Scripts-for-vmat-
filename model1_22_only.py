"""
=============================================================
Model 1b — QA Pass/Fail Prediction, dosisoft 2%/2mm ONLY
Dataset: Per_protocol.xlsx, restricted to the 2%/2 protocol
  PASS (1): ' GPR' sheet, Protocol=dosisoft, Status=ΕΠΙΤΥΧΙΑ
  FAIL (0): ' GPR' sheet, Protocol=dosisoft, Status=ΑΠΟΤΥΧΙΑ
          + ' Fail QA' sheet (failed ALL protocols)

NOTE: N per group is printed at runtime, not hardcoded.

WHY A SEPARATE SCRIPT :
Restricting to a single protocol makes 'Protocol' constant across
every row. A constant column carries zero information for a
classifier — so Protocol_enc is removed from the feature set
entirely here. This forces the model to rely purely on the
complexity metrics + anatomical site, which is the actual point
of this protocol-specific variant.

Features: 12 complexity metrics + anatomical site
Models: Random Forest + XGBoost
=============================================================
"""

# ================================================================
#  IMPORTANT — NOT FOR GENERAL USE
# ================================================================
# This script is provided for reference and reproducibility of the
# results presented in this thesis. It requires the labelled dataset
# 'Per_protocol.xlsx' which contains patient-related clinical data
# and CANNOT be shared due to confidentiality.
#
# This script CANNOT be executed without the corresponding dataset.
# It is NOT intended as a general-purpose tool for other datasets.
#
# For reusable code, see:
#   - all_metrics.py   (VMAT complexity extraction from DICOM-RT)
#   - csv_to_excel.py  (CSV to Excel converter)
# ================================================================

import pandas as pd
import numpy as np
import matplotlib.pyplot as plt
from sklearn.ensemble import RandomForestClassifier
from sklearn.model_selection import StratifiedKFold, cross_validate, cross_val_predict
from sklearn.metrics import roc_curve, auc, confusion_matrix, classification_report
from sklearn.preprocessing import LabelEncoder
import joblib
import openpyxl
import warnings
warnings.filterwarnings('ignore')

try:
    from xgboost import XGBClassifier
    HAS_XGB = True
except ImportError:
    print("  XGBoost not found. Install with: pip install xgboost")
    HAS_XGB = False


# ============================================================
#  USER SETTINGS — CHANGE THIS PATH FOR YOUR SYSTEM
# ============================================================

FILE = "Per_protocol.xlsx"  # Change to full path if needed

# ============================================================


# =============================================================
# 1. DATA LOADING (dosisoft only)
# =============================================================
print("=" * 60)
print(" MODEL 1b — QA PASS/FAIL PREDICTION (dosisoft 2%/2mm ONLY)")
print("=" * 60)

METRIC_COLS = [
    'MCSv', 'Total_MU', 'LSV_Mean', 'AAV_Mean', 'MI_Mean',
    'Leaf_Travel_mm', 'LTMCS', 'MU_per_CP_Mean', 'MU_lt_3_Percent',
    'SAS5_Mean', 'SAS10_Mean', 'AAJA_Mean'
]

def norm_proto(p):
    """Normalize protocol name to 4 categories."""
    if not p or str(p).strip() in ('', 'nan', 'Unknown'): return 'Unknown'
    p = str(p).replace('\xa0', ' ').lower().strip()
    if 'dosisoft' in p: return 'dosisoft'
    if '90%' in p:      return '3mm90'
    if '3% 2mm' in p or '3%2mm' in p: return '3mm2'
    if '3% 3mm' in p or '3%3mm' in p: return '3mm3'
    return 'Unknown'

def norm_site(d):
    """Map diagnosis to anatomical site group."""
    if not d: return 'Unknown'
    d = str(d).upper()
    if 'PROSTATE' in d:                                                return 'Prostate'
    if any(x in d for x in ['ENDOMETRIUM','CERVIX','VULVA','OVARY']):  return 'Gynecological'
    if any(x in d for x in ['BLADDER','RECTUM','ANUS']):               return 'Pelvic'
    return 'Other'

# Robust sheet lookup (handles inconsistent leading/trailing spaces)
wb = openpyxl.load_workbook(FILE, read_only=True)
sheet_map = {s.strip().lower(): s for s in wb.sheetnames}

def find_sheet(target):
    key = target.strip().lower()
    if key not in sheet_map:
        raise ValueError(f"Could not find sheet matching '{target}'. Sheets: {wb.sheetnames}")
    return sheet_map[key]

GPR_SHEET  = find_sheet('GPR')
FAIL_SHEET = find_sheet('Fail QA')

# ── Source 1: GPR sheet, FILTERED to dosisoft protocol only ─
df_gpr = pd.read_excel(FILE, sheet_name=GPR_SHEET)
df_gpr['QA_Label'] = df_gpr['Status'].map({'ΕΠΙΤΥΧΙΑ': 1, 'ΑΠΟΤΥΧΙΑ': 0})
df_gpr['Protocol'] = df_gpr['Πρωτόκολλο'].apply(norm_proto)
df_gpr['Site']     = df_gpr['Διάγνωση'].apply(norm_site)

df_gpr_dosisoft = df_gpr[df_gpr['Protocol'] == 'dosisoft'].copy()
df_gpr_clean = df_gpr_dosisoft[METRIC_COLS + ['Site', 'QA_Label']].dropna(subset=['QA_Label'])

# ── Source 2: Fail QA sheet -- ALL rows count as dosisoft FAIL ──
# (dosisoft is always tried first at this center; every plan here failed it)
df_fail = pd.read_excel(FILE, sheet_name=FAIL_SHEET)
df_fail['Site']     = df_fail['Διάγνωση (filename)'].apply(norm_site)
df_fail['QA_Label'] = 0  # FAIL
df_fail_clean = df_fail[METRIC_COLS + ['Site', 'QA_Label']]

# ── Combine ──────────────────────────────────────────────────
df = pd.concat([df_gpr_clean, df_fail_clean], ignore_index=True)
df = df.dropna(subset=METRIC_COLS)

n_pass = int((df['QA_Label'] == 1).sum())
n_fail = len(df) - n_pass

print(f"\nDataset composition (dosisoft 2%/2mm only):")
print(f"  PASS: {n_pass}")
print(f"  FAIL: {n_fail}")
print(f"  TOTAL: {len(df)}  |  Ratio PASS:FAIL = {n_pass/n_fail:.2f}:1")

# ── Encode site ONLY (protocol removed -- constant, zero information) ──
le_site = LabelEncoder()
df['Site_enc'] = le_site.fit_transform(df['Site'].astype(str))

print(f"\n  Sites: {list(le_site.classes_)}")
print(f"\n  NOTE: Protocol_enc is NOT a feature here -- every row is dosisoft,")
print(f"  so protocol carries zero information for this model by design.")

FEATURE_COLS = METRIC_COLS + ['Site_enc']
X = df[FEATURE_COLS].fillna(df[FEATURE_COLS].median(numeric_only=True))
y = df['QA_Label']


# =============================================================
# 2. MODEL DEFINITIONS
# =============================================================
print("\n" + "=" * 60)
print(" MODEL DEFINITIONS")
print("=" * 60)

models = {
    "Random Forest": RandomForestClassifier(
        n_estimators=300,
        class_weight='balanced',
        min_samples_split=5,
        min_samples_leaf=2,
        max_features='sqrt',
        random_state=42,
        n_jobs=-1
    ),
}
if HAS_XGB:
    models["XGBoost"] = XGBClassifier(
        n_estimators=300,
        learning_rate=0.05,
        max_depth=4,
        subsample=0.8,
        colsample_bytree=0.8,
        objective='binary:logistic',
        random_state=42,
        verbosity=0,
        eval_metric='logloss'
    )

for name in models:
    print(f"  ✓ {name}")


# =============================================================
# 3. STRATIFIED 5-FOLD CROSS VALIDATION
# =============================================================
print("\n" + "=" * 60)
print(" STRATIFIED 5-FOLD CROSS VALIDATION")
print("=" * 60)

skf = StratifiedKFold(n_splits=5, shuffle=True, random_state=42)
results = {}
roc_data = {}

for name, model in models.items():
    cv = cross_validate(
        model, X, y, cv=skf,
        scoring=['roc_auc', 'accuracy', 'f1', 'precision', 'recall'],
        return_train_score=False
    )
    results[name] = cv

    y_proba = cross_val_predict(model, X, y, cv=skf, method='predict_proba')[:, 1]
    fpr, tpr, thresholds = roc_curve(y, y_proba)
    roc_auc = auc(fpr, tpr)
    roc_data[name] = (fpr, tpr, roc_auc, thresholds)

    # Optimal threshold via Youden's J statistic
    j_scores = tpr - fpr
    best_idx = np.argmax(j_scores)
    best_threshold = thresholds[best_idx]

    print(f"\n{'─'*50}")
    print(f" {name}")
    print(f"{'─'*50}")
    print(f"  AUC       : {cv['test_roc_auc'].mean():.3f} ± {cv['test_roc_auc'].std():.3f}")
    print(f"  Accuracy  : {cv['test_accuracy'].mean():.3f} ± {cv['test_accuracy'].std():.3f}")
    print(f"  F1        : {cv['test_f1'].mean():.3f} ± {cv['test_f1'].std():.3f}")
    print(f"  Precision : {cv['test_precision'].mean():.3f} ± {cv['test_precision'].std():.3f}")
    print(f"  Recall    : {cv['test_recall'].mean():.3f} ± {cv['test_recall'].std():.3f}")
    print(f"  Optimal threshold (Youden's J): {best_threshold:.3f}")

    y_pred_opt = (y_proba >= best_threshold).astype(int)
    print(f"\n  Classification report (pooled OOF predictions, optimal threshold):")
    print(classification_report(y, y_pred_opt, target_names=['FAIL', 'PASS'], digits=3))


# =============================================================
# 4. COMPARISON TABLE
# =============================================================
print("\n" + "=" * 60)
print(" RESULTS COMPARISON TABLE")
print("=" * 60)

summary = pd.DataFrame({
    name: {
        "AUC":        f"{v['test_roc_auc'].mean():.3f}",
        "AUC (±std)": f"±{v['test_roc_auc'].std():.3f}",
        "Accuracy":   f"{v['test_accuracy'].mean():.3f}",
        "F1":         f"{v['test_f1'].mean():.3f}",
        "Precision":  f"{v['test_precision'].mean():.3f}",
        "Recall":     f"{v['test_recall'].mean():.3f}",
    }
    for name, v in results.items()
}).T
print(summary.to_string())
best = max(results, key=lambda n: results[n]['test_roc_auc'].mean())
print(f"\n Best model (AUC): {best}")

print("\nComparison with the full mixed-protocol Model 1:")
print("this variant removes Protocol_enc entirely, so any AUC achieved here")
print("comes purely from the 12 complexity metrics + site -- a cleaner test")
print("of whether complexity alone predicts dosisoft-specific outcome.")


# =============================================================
# 5. FEATURE IMPORTANCE + SAVE MODELS
# =============================================================
print("\n" + "=" * 60)
print(" FEATURE IMPORTANCE (full dataset training)")
print("=" * 60)

importances = {}
trained_models = {}

rf_full = RandomForestClassifier(n_estimators=300, class_weight='balanced',
                                  min_samples_split=5, min_samples_leaf=2,
                                  max_features='sqrt', random_state=42, n_jobs=-1)
rf_full.fit(X, y)
importances["Random Forest"] = pd.Series(rf_full.feature_importances_,
                                          index=FEATURE_COLS).sort_values(ascending=False)
trained_models["Random Forest"] = rf_full

if HAS_XGB:
    xgb_full = XGBClassifier(n_estimators=300, learning_rate=0.05, max_depth=4,
                              subsample=0.8, colsample_bytree=0.8,
                              objective='binary:logistic',
                              random_state=42, verbosity=0, eval_metric='logloss')
    xgb_full.fit(X, y)
    importances["XGBoost"] = pd.Series(xgb_full.feature_importances_,
                                        index=FEATURE_COLS).sort_values(ascending=False)
    trained_models["XGBoost"] = xgb_full

for mname, imp in importances.items():
    print(f"\n{mname} — Top 8:")
    for i, (feat, val) in enumerate(imp.head(8).items(), 1):
        print(f"  {i}. {feat:<25} {val:.4f}")

for mname, model in trained_models.items():
    fname = f"model1b_{'RF' if 'Forest' in mname else 'XGB'}.pkl"
    joblib.dump({
        'model': model,
        'feature_cols': FEATURE_COLS,
        'metric_cols': METRIC_COLS,
        'le_site': le_site,
        'threshold': 0.5
    }, fname)
    print(f"\n  ✓ Saved: {fname}")


# =============================================================
# 6. PLOTS
# =============================================================
print("\n" + "=" * 60)
print(" GENERATING PLOTS")
print("=" * 60)

colors = ["#55A868", "#C44E52", "#4C72B0"]

fig, axes = plt.subplots(2, 2, figsize=(14, 12))
fig.suptitle(f"Model 1b — QA Pass/Fail Prediction, dosisoft 2%/2mm only\n"
             f"(RF + XGBoost, N={len(df)}: {n_pass} PASS + {n_fail} FAIL, no Protocol feature)",
             fontsize=13, fontweight='bold')

# ROC Curves
ax1 = axes[0, 0]
for name, color in zip(roc_data, colors):
    fpr, tpr, roc_auc, _ = roc_data[name]
    ax1.plot(fpr, tpr, color=color, lw=2, label=f'{name} (AUC={roc_auc:.3f})')
ax1.plot([0, 1], [0, 1], 'k--', lw=1, label='Random classifier (AUC=0.5)')
ax1.set_xlabel("1 - Specificity (False Positive Rate)")
ax1.set_ylabel("Sensitivity (True Positive Rate)")
ax1.set_title("ROC Curves — dosisoft only")
ax1.legend(fontsize=8, loc='lower right')

# AUC per fold
ax2 = axes[0, 1]
auc_data = [results[n]['test_roc_auc'] for n in results]
bp = ax2.boxplot(auc_data, patch_artist=True, widths=0.5)
for patch, color in zip(bp['boxes'], colors[:len(results)]):
    patch.set_facecolor(color)
    patch.set_alpha(0.75)
ax2.set_xticks(range(1, len(results)+1))
ax2.set_xticklabels(list(results.keys()), rotation=15, ha='right', fontsize=9)
ax2.set_ylabel("AUC")
ax2.set_title("AUC per Fold (5-CV)")
ax2.axhline(0.5, color='gray', linestyle='--', alpha=0.5)
ax2.set_ylim(0, 1)

# Feature importance
ax3 = axes[1, 0]
top_feats = importances["Random Forest"].head(10)
bars = ax3.barh(top_feats.index[::-1], top_feats.values[::-1],
                color="#55A868", alpha=0.8, edgecolor='white')
ax3.set_xlabel("Feature Importance")
ax3.set_title("Top 10 Features — Random Forest")
for bar, val in zip(bars, top_feats.values[::-1]):
    ax3.text(bar.get_width()+0.002, bar.get_y()+bar.get_height()/2,
             f'{val:.3f}', va='center', fontsize=9)

# Dataset composition
ax4 = axes[1, 1]
n_fail_proto = (df_gpr_clean['QA_Label'] == 0).sum()
n_fail_all = len(df_fail_clean)
labels = [f'PASS\n({n_pass})', f'FAIL - dosisoft\n({n_fail_proto})', f'FAIL - all\n({n_fail_all})']
sizes  = [n_pass, n_fail_proto, n_fail_all]
clrs   = ['#55A868', '#FF8C00', '#C44E52']
bars4 = ax4.bar(labels, sizes, color=clrs, alpha=0.8, edgecolor='white')
for bar, val in zip(bars4, sizes):
    ax4.text(bar.get_x()+bar.get_width()/2, bar.get_height()+1,
             str(val), ha='center', fontweight='bold')
ax4.set_ylabel("Number of plans")
ax4.set_title(f"Dataset Composition\n(Total N={len(df)})")
ax4.set_ylim(0, max(sizes) * 1.2)

plt.tight_layout()
plt.savefig("model1b_dosisoft_only.png", dpi=150, bbox_inches='tight')
plt.show()
print("\n✅ Plot saved as 'model1b_dosisoft_only.png'")
print("✅ Model 1b complete!")